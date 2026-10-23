#!/usr/bin/env python3
"""
05_planilhas.py — Gera a planilha editável planilhas/calculos_papeis_de_trabalho.xlsx
(com FÓRMULAS vivas): materialidade, conciliações, analíticos, liquidez/continuidade, partes relacionadas,
parâmetros da amostra, mapa de distorções e comparações com a base sintética.
Valores de entrada (azul) vêm das DFs públicas de 31/12/2025; resultados são fórmulas.
Classificação: itens baseados só nas DFs = "Realizado com evidência pública"; itens da base = "Simulado para fins didáticos".
"""
import os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
S = os.path.join(RAIZ, "saidas")
OUT = os.path.join(RAIZ, "planilhas", "calculos_papeis_de_trabalho.xlsx")

AZUL = Font(color="0000FF"); NEG = Font(bold=True); HDR = PatternFill("solid", fgColor="DDE6F0")
wb = Workbook()
NUM = '#,##0;[Red](#,##0)'; PCT = '0.0%'; DEC = '#,##0.00;[Red](#,##0.00)'


def cab(ws, row, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v); c.font = NEG; c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def larg(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================ Dados_DF ============================
ws = wb.active; ws.title = "Dados_DF"
ws["A1"] = "DADOS PÚBLICOS — Bradespar S.A. — DFs de 31/12/2025 (R$ mil, salvo indicação). Azul = entrada transcrita da fonte."
ws["A1"].font = NEG
cab(ws, 3, ["Chave", "Descrição", "2025", "2024", "Fonte (documento / página impressa)"])
DADOS = [
    ("ativo_total", "Total do ativo", 7999560, 8606884, "Demonstrações Financeiras (DF) p.14 — Balanço"),
    ("ativo_circ", "Ativo circulante", 565128, 360610, "DF p.14"),
    ("caixa", "Caixa e equivalentes de caixa", 18402, 287950, "DF p.14; Nota 7 p.26"),
    ("div_receber", "Dividendos e JCP a receber", 546640, 72276, "DF p.14"),
    ("investimentos", "Investimentos (Vale — equivalência)", 7344756, 8153681, "DF p.14; Nota 8 p.27"),
    ("imobilizado", "Imobilizado", 899, 1011, "DF p.14"),
    ("intangivel", "Intangível", 4067, 2306, "DF p.14; Nota 16b p.34"),
    ("tributos_rlp", "Tributos a compensar (RLP)", 84710, 89276, "DF p.14; Nota 12a p.31"),
    ("passivo_circ", "Passivo circulante", 347821, 121412, "DF p.14"),
    ("passivo_nc", "Passivo não circulante", 3968, 2105, "DF p.14"),
    ("pl", "Patrimônio líquido", 7647771, 8483367, "DF p.14"),
    ("div_pagar", "Dividendos e JCP a pagar", 316563, 87936, "DF p.14"),
    ("outras_obrig", "Outras obrigações", 30034, 32545, "DF p.14; Nota 16c p.34"),
    ("lucro", "Lucro líquido do exercício", 546993, 1192208, "DF p.15 — DRE"),
    ("eq_pat", "Resultado de equivalência patrimonial", 571264, 1187822, "DF p.15; Nota 8 p.27"),
    ("dga", "Despesas gerais e administrativas", -34541, -26286, "DF p.15; Nota 16d p.34"),
    ("desp_trib", "Despesas tributárias", -64815, -42823, "DF p.15"),
    ("rec_fin", "Receitas financeiras", 75615, 58167, "DF p.15; Nota 11 p.31"),
    ("rec_aplic", "  Receita de aplicações financeiras", 55506, 42312, "Nota 11 p.31"),
    ("res_fin", "Resultado financeiro", 75085, 57639, "DF p.15"),
    ("oci", "Reflexos de empresas com influência significativa (OCI)", -135589, 329140, "DF p.16 — DRA; DMPL p.17"),
    ("dfc_op", "Caixa líquido das atividades operacionais", 670015, 790041, "DF p.18 — DFC"),
    ("dfc_fin", "Caixa líquido das atividades de financiamento", -939563, -747196, "DF p.18"),
    ("dfc_receb", "JCP e dividendos recebidos (DFC)", 685334, 738081, "DF p.18"),
    ("var_caixa", "Variação de caixa", -269548, 41746, "DF p.18"),
    ("jcp_receb", "JCP recebidos de Vale (bruto IRRF)", 309620, 342032, "Nota 13 I p.32"),
    ("jcp_receber", "JCP a receber de Vale (bruto IRRF)", 256391, 85031, "Nota 13 I p.32"),
    ("div_receb", "Dividendos recebidos de Vale", 349881, 447354, "Nota 13 I p.32"),
    ("div_receber_v", "Dividendos a receber de Vale", 328708, None, "Nota 13 I p.32 (2024 não informado)"),
    ("irrf", "Alíquota de IRRF sobre JCP", 0.15, 0.15, "Nota 10c p.29–30"),
    ("vale_pl_aj", "Vale — PL ajustado (Nota 8)", 204083020, None, "Nota 8 p.27"),
    ("vale_res_aj", "Vale — resultado ajustado (Nota 8)", 13814281, 31592000, "Nota 8 p.27; Nota 18 p.35 (2024 em R$ mi)"),
    ("p_dir", "Participação no capital social (Nota 8)", 0.0360, 0.0360, "Nota 8 p.27; Nota 18 p.35"),
    ("p_ex_tes", "Participação sem ações em tesouraria (V/T)", 0.0383, None, "Relatório da Administração — Estrutura dos Investimentos, p.3"),
    ("mercado", "Valor de mercado da participação na Vale", 11754997, 8910993, "Nota 8 (3) p.27"),
    ("div_dec_total", "Dividendos declarados 2025 pela Bradespar (DMPL)", 680000, 450000, "DMPL p.17"),
    ("jcp_dec_total", "JCP declarados 2025 pela Bradespar (DMPL)", 567000, 422000, "DMPL p.17"),
    ("rem_adm", "Remuneração do pessoal-chave (total)", 18660, 15000, "Nota 13 II p.32"),
    ("rem_limite", "Limite AGO 2025: remuneração + previdência", 19910, None, "Nota 13 II p.32 (19.391 + 519)"),
    ("desp_pessoal", "Despesas de pessoal", 26755, 18463, "Nota 16d p.34"),
    ("outras_dga", "Outras despesas gerais e administrativas", 7786, 7823, "Nota 16d p.34"),
    ("selic_proxy", "Taxa prefixada 1 ano usada no cenário-base da Nota 14", 0.138, None, "Nota 14b p.33–34 (exemplo 13,80% a.a.)"),
    ("litel", "Contingência possível — Litel (R$ mil)", 1400000, None, "Nota 9 b-I p.27"),
    ("autuacao", "Contingência possível — autuação IRPJ/CSL (R$ mil)", 979000, None, "Nota 9 b-II p.28"),
    ("cred_trib", "Créditos tributários não ativados", 355800, 345171, "Nota 12c p.32"),
]
R = {}
for i, (k, d, a, b, f) in enumerate(DADOS, start=4):
    ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=d)
    for col, v in ((3, a), (4, b)):
        c = ws.cell(row=i, column=col, value=v); c.font = AZUL
        c.number_format = PCT if isinstance(v, float) and abs(v) < 1 else NUM
    ws.cell(row=i, column=5, value=f); R[k] = i
larg(ws, [16, 58, 16, 16, 70])


def d(k, ano=2025):
    return f"Dados_DF!${'C' if ano == 2025 else 'D'}${R[k]}"


# ============================ Materialidade ============================
ws = wb.create_sheet("Materialidade")
ws["A1"] = "PT 320 — Materialidade (R$ mil) — decisão didática; julgamento do preparador (matrícula 211014396)"; ws["A1"].font = NEG
cab(ws, 3, ["Referencial", "Base (R$ mil)", "% aplicado", "Resultado (R$ mil)", "Comentário"])
refs = [
    ("Ativo total", f"={d('ativo_total')}", 0.01, "ESCOLHIDO para as DFs como um todo: ativo concentrado (92% em Investimentos) e usuários focados no valor da participação na Vale."),
    ("Patrimônio líquido", f"={d('pl')}", 0.01, "Alternativa convergente."),
    ("Lucro antes do IR (resultado)", f"={d('lucro')}", 0.05, "Volátil (−54% a/a) e derivado da investida: usado para MATERIALIDADE ESPECÍFICA de contas de resultado."),
    ("Equivalência patrimonial (proxy de receita)", f"={d('eq_pat')}", 0.05, "Referencial alternativo de resultado."),
]
for i, (n, f, p, c) in enumerate(refs, start=4):
    ws.cell(row=i, column=1, value=n); ws.cell(row=i, column=2, value=f).number_format = NUM
    x = ws.cell(row=i, column=3, value=p); x.font = AZUL; x.number_format = PCT
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").number_format = NUM; ws.cell(row=i, column=5, value=c)
ws["A9"] = "Materialidade global adotada (arredondada)"; ws["A9"].font = NEG
ws["D9"] = "=ROUND(D4,-3)"; ws["D9"].number_format = NUM
ws["A10"] = "Percentual para execução"; ws["C10"] = 0.70; ws["C10"].font = AZUL; ws["C10"].number_format = PCT
ws["A11"] = "Materialidade para execução (PM)"; ws["D11"] = "=D9*C10"; ws["D11"].number_format = NUM
ws["A12"] = "Limite de claramente trivial"; ws["C12"] = 0.05; ws["C12"].font = AZUL; ws["C12"].number_format = PCT
ws["D12"] = "=D9*C12"; ws["D12"].number_format = NUM
ws["A13"] = "Materialidade específica — contas de resultado e distribuições (5% do lucro)"; ws["D13"] = "=ROUND(D6,-1)"; ws["D13"].number_format = NUM
ws["A14"] = "PM específica (70%)"; ws["D14"] = "=D13*C10"; ws["D14"].number_format = NUM
ws["A15"] = "Materialidade específica — remuneração de administradores / partes relacionadas"; ws["C15"] = 0.05; ws["C15"].font = AZUL; ws["C15"].number_format = PCT
ws["D15"] = f"=ROUND({d('rem_adm')}*C15,0)"; ws["D15"].number_format = NUM
ws["E15"] = "Sensibilidade qualitativa (NBC TA 320.A..; 550): 5% do total pago ao pessoal-chave."
ws["A17"] = "Teste de razoabilidade"; ws["A17"].font = NEG
ws["A18"] = "Materialidade global / lucro 2025"; ws["D18"] = f"=D9/{d('lucro')}"; ws["D18"].number_format = PCT
ws["A19"] = "Materialidade global / PL"; ws["D19"] = f"=D9/{d('pl')}"; ws["D19"].number_format = PCT
ws["A20"] = "Materialidade global / valor de mercado da Vale"; ws["D20"] = f"=D9/{d('mercado')}"; ws["D20"].number_format = PCT
ws["A21"] = "Investimentos / ativo total"; ws["D21"] = f"={d('investimentos')}/{d('ativo_total')}"; ws["D21"].number_format = PCT
larg(ws, [70, 16, 12, 18, 100])

# ============================ Concil_Investimento ============================
ws = wb.create_sheet("Concil_Investimento")
ws["A1"] = "PT 500/510/330 — Roll-forward da conta Investimentos e conciliações com Notas 8, 10 e 13 (R$ mil)"; ws["A1"].font = NEG
cab(ws, 3, ["Item", "2025", "Fonte/observação"])
rf = [
    ("Saldo inicial (31/12/2024)", f"={d('investimentos',2024)}", "Balanço 2024 (comparativo)"),
    ("(+) Equivalência patrimonial", f"={d('eq_pat')}", "DRE; Nota 8"),
    ("(+/−) Reflexos em OCI", f"={d('oci')}", "DRA/DMPL"),
    ("(−) JCP declarados pela Vale (bruto): recebidos + a receber", f"=-({d('jcp_receb')}+{d('jcp_receber')})", "Nota 13 I"),
    ("(−) Dividendos declarados pela Vale: recebidos + a receber", f"=-({d('div_receb')}+{d('div_receber_v')})", "Nota 13 I"),
    ("Total: saldo final calculado", "=SUM(B4:B8)", ""),
    ("Saldo final no Balanço", f"={d('investimentos')}", "Balanço"),
    ("Diferença (deve ser zero)", "=B9-B10", "Conciliação exata em R$ mil"),
]
for i, (a, b, c) in enumerate(rf, start=4):
    ws.cell(row=i, column=1, value=a); ws.cell(row=i, column=2, value=b).number_format = NUM; ws.cell(row=i, column=3, value=c)
ws["A13"] = "Conciliação de dividendos e JCP a receber (Balanço × Nota 13)"; ws["A13"].font = NEG
cr = [
    ("JCP a receber bruto 2025", f"={d('jcp_receber')}", ""),
    ("(−) IRRF 15%", f"=-B14*{d('irrf')}", "Nota 10c: IRRF 15% sobre JCP"),
    ("(+) Dividendos a receber", f"={d('div_receber_v')}", ""),
    ("Total: dividendos e JCP a receber calculado", "=SUM(B14:B16)", ""),
    ("Balanço (Dividendos e JCP a receber)", f"={d('div_receber')}", ""),
    ("Diferença (arredondamento)", "=B17-B18", "Aceitável: < R$ 1 mil"),
    ("2024: JCP a receber bruto − IRRF", f"={d('jcp_receber',2024)}*(1-{d('irrf',2024)})", ""),
    ("2024: Balanço", f"={d('div_receber',2024)}", ""),
    ("2024: diferença", "=B20-B21", ""),
]
for i, (a, b, c) in enumerate(cr, start=14):
    ws.cell(row=i, column=1, value=a); ws.cell(row=i, column=2, value=b).number_format = DEC; ws.cell(row=i, column=3, value=c)
ws["A24"] = "Conciliação do fluxo de caixa: JCP e dividendos recebidos (DFC) × Nota 13"; ws["A24"].font = NEG
cf = [
    ("JCP recebidos 2025 (bruto)", f"={d('jcp_receb')}"),
    ("(−) IRRF 15% retido", f"=-B25*{d('irrf')}"),
    ("(+) Dividendos recebidos 2025", f"={d('div_receb')}"),
    ("(+) Recebimento em 2025 do saldo a receber de 2024", f"={d('div_receber',2024)}"),
    ("Total: recebimentos calculados", "=SUM(B25:B28)"),
    ("DFC — JCP e dividendos recebidos", f"={d('dfc_receb')}"),
    ("Diferença (deve ser ~0)", "=B29-B30"),
]
for i, (a, b) in enumerate(cf, start=25):
    ws.cell(row=i, column=1, value=a); ws.cell(row=i, column=2, value=b).number_format = DEC
ws["A33"] = "Verificação do valor patrimonial × participação"; ws["A33"].font = NEG
vp = [
    ("Participação (Nota 8) × PL ajustado da Vale", f"={d('p_dir')}*{d('vale_pl_aj')}"),
    ("Saldo contábil", f"={d('investimentos')}"),
    ("Diferença", "=B35-B34"),
    ("Diferença / saldo", "=B36/B35"),
    ("Valor de mercado da participação", f"={d('mercado')}"),
    ("Mercado / contábil (indício de impairment? <1 = sim)", "=B38/B35"),
]
for i, (a, b) in enumerate(vp, start=34):
    c = ws.cell(row=i, column=1, value=a); x = ws.cell(row=i, column=2, value=b); x.number_format = PCT if i == 37 else (DEC if i != 39 else '0.00')
larg(ws, [75, 18, 50])

# ============================ Analiticos ============================
ws = wb.create_sheet("Analiticos")
ws["A1"] = "PT 520 — Procedimentos analíticos substantivos (R$ mil) — Realizado com evidência pública"; ws["A1"].font = NEG
ws["A2"] = "Limite de diferença aceitável sem investigação (NBC TA 520.5(d)) = PM específica de resultado (Materialidade!D14)"
cab(ws, 4, ["Análise", "Expectativa", "Registrado", "Diferença", "Dif./registrado", "Limite aceitável", "Investigar?", "Comentário"])
an = [
    ("A1 Equivalência: 3,60% × resultado ajustado da Vale", f"={d('p_dir')}*{d('vale_res_aj')}", f"={d('eq_pat')}",
     "Participação direta sobre capital total (inclui ações em tesouraria da investida)."),
    ("A2 Equivalência: 3,83% (sem tesouraria) × resultado ajustado", f"={d('p_ex_tes')}*{d('vale_res_aj')}", f"={d('eq_pat')}",
     "Expectativa mais precisa; resíduo ainda acima do limite: a % média ponderada ao longo de 2025 pode ser maior (recompras) — requer ITRs da Vale."),
    ("A3 Investimento: 3,60% × PL ajustado da Vale", f"={d('p_dir')}*{d('vale_pl_aj')}", f"={d('investimentos')}", "Alta aderência (0,03%)."),
    ("A4 Receita de aplicações: taxa 13,80% × média simples do caixa", f"={d('selic_proxy')}*AVERAGE({d('caixa')},{d('caixa',2024)})", f"={d('rec_aplic')}",
     "Média simples não reflete saldos intra-anuais (JCP/dividendos recebidos de R$ 685 mi e pagamentos de R$ 940 mi). Precisão insuficiente."),
    ("A5 Despesas G&A: pessoal 2025 + outras despesas de 2024", f"={d('desp_pessoal')}+{d('outras_dga',2024)}", f"=-{d('dga')}",
     "Aumento de G&A explicado pelo aumento de pessoal (R$ 18.463 → 26.755 mil)."),
    ("A6 Remuneração do pessoal-chave vs limite aprovado em AGO", f"={d('rem_limite')}", f"={d('rem_adm')}", "Dentro do limite (R$ 19.391 + R$ 519)."),
    ("A7 DFC: variação de caixa = op. + invest. + financ.", f"={d('dfc_op')}+{d('dfc_fin')}", f"={d('var_caixa')}", "Fecha com 0 de investimentos em 2025."),
    ("A8 Caixa final = inicial + variação", f"={d('caixa',2024)}+{d('var_caixa')}", f"={d('caixa')}", "Fecha exatamente."),
    ("A9 Dividendos+JCP declarados (DMPL) × declarados Bradespar (Nota 10c)", "=1247000", f"={d('div_dec_total')}+{d('jcp_dec_total')}", "Nota 10c total 2025 R$ 1.247.000 mil vs DMPL R$ 1.247.000 mil."),
]
for i, (a, e, r, c) in enumerate(an, start=5):
    ws.cell(row=i, column=1, value=a)
    ws.cell(row=i, column=2, value=e).number_format = NUM; ws.cell(row=i, column=3, value=r).number_format = NUM
    ws.cell(row=i, column=4, value=f"=C{i}-B{i}").number_format = NUM
    ws.cell(row=i, column=5, value=f"=IF(C{i}=0,0,D{i}/C{i})").number_format = PCT
    ws.cell(row=i, column=6, value="=Materialidade!$D$14").number_format = NUM
    ws.cell(row=i, column=7, value=f'=IF(ABS(D{i})>F{i},"SIM","não")')
    ws.cell(row=i, column=8, value=c)
larg(ws, [66, 16, 16, 16, 14, 16, 12, 100])

# ============================ Liquidez_Continuidade ============================
ws = wb.create_sheet("Liquidez_Continuidade")
ws["A1"] = "PT 570 — Continuidade operacional / liquidez (R$ mil) — Realizado com evidência pública (limitado às DFs)"; ws["A1"].font = NEG
cab(ws, 3, ["Indicador", "2025", "2024", "Comentário"])
li = [
    ("Liquidez corrente (AC/PC)", f"={d('ativo_circ')}/{d('passivo_circ')}", f"={d('ativo_circ',2024)}/{d('passivo_circ',2024)}", "Reduziu em 2025, mas permanece > 1", '0.00'),
    ("Caixa / passivo circulante", f"={d('caixa')}/{d('passivo_circ')}", f"={d('caixa',2024)}/{d('passivo_circ',2024)}", "Cobertura por caixa caiu: depende de recebimentos da Vale", '0.00'),
    ("(Caixa + Div/JCP a receber) / Div/JCP a pagar", f"=({d('caixa')}+{d('div_receber')})/{d('div_pagar')}", f"=({d('caixa',2024)}+{d('div_receber',2024)})/{d('div_pagar',2024)}", "Obrigação com acionistas coberta por recebíveis da Vale", '0.00'),
    ("Dívida financeira", 0, 0, "Sem empréstimos/financiamentos nas DFs (Notas 14 e 16)", NUM),
    ("Caixa operacional / distribuições pagas", f"={d('dfc_op')}/-{d('dfc_fin')}", f"={d('dfc_op',2024)}/-{d('dfc_fin',2024)}", "<1: distribuições superiores ao caixa gerado; consumo do caixa acumulado", '0.00'),
    ("Passivos contingentes possíveis (Litel + autuação)", f"={d('litel')}+{d('autuacao')}", None, "Não provisionados (Nota 9)", NUM),
    ("  … como % do PL", f"=B9/{d('pl')}", None, "", PCT),
    ("  … como % do valor de mercado da Vale", f"=B9/{d('mercado')}", None, "Capacidade de pagamento via venda de ações, em tese", PCT),
    ("  … vezes o caixa e equivalentes", f"=B9/{d('caixa')}", None, "Sem liquidez imediata — gatilho de atenção se houver perda provável", '0.0'),
    ("Créditos tributários não ativados", f"={d('cred_trib')}", f"={d('cred_trib',2024)}", "Não reconhecidos; sem impacto no balanço", NUM),
]
for i, (a, b, c, com, fmt) in enumerate(li, start=4):
    ws.cell(row=i, column=1, value=a); x = ws.cell(row=i, column=2, value=b); x.number_format = fmt
    if c is not None:
        y = ws.cell(row=i, column=3, value=c); y.number_format = fmt
    ws.cell(row=i, column=4, value=com)
larg(ws, [60, 16, 16, 90])

# ============================ Amostra_MUS ============================
ws = wb.create_sheet("Parametros_MUS")
ws["A1"] = "PT 530 — Parâmetros da amostragem por unidade monetária (R$) — BASE SINTÉTICA — Simulado para fins didáticos"; ws["A1"].font = NEG
mus = pd.read_csv(os.path.join(S, "mus_resumo.csv"))
vals = dict(zip(mus["parametro_ou_resultado"], mus["valor"]))
ws["A3"] = "Valor contábil da população (R$)"; ws["B3"] = float(vals["Valor contábil da população (R$)"]); ws["B3"].font = AZUL
ws["A4"] = "Distorção tolerável = PM (R$)"; ws["B4"] = "=Materialidade!D11*1000"
ws["A5"] = "Distorção esperada (R$) — hipótese"; ws["B5"] = 5000000; ws["B5"].font = AZUL
ws["A6"] = "Fator de confiança (zero erros, 95%)"; ws["B6"] = 2.9957; ws["B6"].font = AZUL
ws["A7"] = "Fator de expansão (95%)"; ws["B7"] = 1.6; ws["B7"].font = AZUL
ws["A8"] = "Tamanho da amostra = VC × FC / (TM − DE × FE)"; ws["B8"] = "=ROUNDUP(B3*B6/(B4-B5*B7),0)"
ws["A9"] = "Intervalo de amostragem (R$)"; ws["B9"] = "=B3/B8"
for r in (3, 4, 5, 9):
    ws[f"B{r}"].number_format = NUM
r0 = 11
ws.cell(row=r0, column=1, value="Resultados da execução (script 03_amostragem_mus.py)").font = NEG
for i, (k, v) in enumerate(vals.items(), start=r0 + 1):
    try:
        v = float(v)
    except (TypeError, ValueError):
        pass
    ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v)
larg(ws, [70, 24])

# ============================ Mapa_Distorcoes ============================
ws = wb.create_sheet("Mapa_Distorcoes")
ws["A1"] = "PT 450 — Mapa de distorções (BASE SINTÉTICA) — valores em R$ — Simulado para fins didáticos"; ws["A1"].font = NEG
cmp_ = pd.read_csv(os.path.join(S, "comparacao_inserido_vs_detectado.csv"))
cab(ws, 3, list(cmp_.columns))
for i, row in enumerate(cmp_.itertuples(index=False), start=4):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        if isinstance(v, float): c.number_format = DEC
n = len(cmp_) + 4
ws.cell(row=n + 1, column=1, value="Materialidade global (R$)"); ws.cell(row=n + 1, column=4, value="=Materialidade!D9*1000").number_format = NUM
ws.cell(row=n + 2, column=1, value="Materialidade para execução (R$)"); ws.cell(row=n + 2, column=4, value="=Materialidade!D11*1000").number_format = NUM
ws.cell(row=n + 3, column=1, value="Limite de trivial (R$)"); ws.cell(row=n + 3, column=4, value="=Materialidade!D12*1000").number_format = NUM
ws.cell(row=n + 4, column=1, value="Efeito líquido (natureza 'valor') / materialidade global")
cmp_val = pd.read_csv(os.path.join(S, "comparacao_resumo.csv"), index_col=0)["valor"]
ws.cell(row=n + 5, column=1, value="Efeito líquido inserido (natureza valor) — R$"); c = ws.cell(row=n + 5, column=4, value=float(cmp_val["efeito_liquido_inserido_valor"])); c.number_format = DEC; c.font = AZUL
ws.cell(row=n + 4, column=4, value=f"=D{n+5}/D{n+1}").number_format = PCT
ws.cell(row=n + 6, column=1, value="Efeito líquido / materialidade para execução"); ws.cell(row=n + 6, column=4, value=f"=D{n+5}/D{n+2}").number_format = PCT
ws.cell(row=n + 7, column=1, value="Parcela que afeta resultado (E1a + E5) — R$")
r_e1a = cmp_[cmp_.tipo_erro == "E1a"]["efeito_inserido_brl"].iloc[0]; r_e5 = cmp_[cmp_.tipo_erro == "E5"]["efeito_inserido_brl"].iloc[0]
c = ws.cell(row=n + 7, column=4, value=float(r_e1a + r_e5)); c.number_format = DEC; c.font = AZUL
ws.cell(row=n + 8, column=1, value="Parcela de resultado / materialidade específica de resultado")
ws.cell(row=n + 8, column=4, value=f"=D{n+7}/(Materialidade!D13*1000)").number_format = PCT
larg(ws, [48, 12, 14, 20, 22, 22, 20, 18, 14])

# ============================ Conc_Razao_Balanco ============================
ws = wb.create_sheet("Conc_Razao_Balanco")
ws["A1"] = "PT 500 — Conciliação do razão auxiliar sintético com o Balanço (R$) — Simulado para fins didáticos"; ws["A1"].font = NEG
conc = pd.read_csv(os.path.join(S, "conciliacao_razao_balanco.csv"))
cab(ws, 3, ["Item", "Valor (R$)"])
for i, row in enumerate(conc.itertuples(index=False), start=4):
    ws.cell(row=i, column=1, value=row[0]); ws.cell(row=i, column=2, value=row[1]).number_format = DEC
larg(ws, [110, 24])

# ============================ Testes_Populacao ============================
ws = wb.create_sheet("Testes_Populacao")
ws["A1"] = "PT 240/520/530 — Resultado dos testes em 100% da base sintética — Simulado para fins didáticos"; ws["A1"].font = NEG
tp = pd.read_csv(os.path.join(S, "resumo_testes_populacao.csv"))
cab(ws, 3, list(tp.columns))
for i, row in enumerate(tp.itertuples(index=False), start=4):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        if isinstance(v, float): c.number_format = DEC
larg(ws, [34, 20, 28, 24])

wb.save(OUT)
print("ok", OUT)
